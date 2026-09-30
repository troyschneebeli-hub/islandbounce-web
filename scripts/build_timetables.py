#!/usr/bin/env python3
"""
Builds data/timetables.js from the operator research spreadsheet, so the
website's real timetables come straight from the sheet (no hand-copying).

    python3 scripts/build_timetables.py path/to/clankinc-operator-research-template.xlsx

Two sheets are read:
  - "Direct Sailings": one row = one published sailing (e.g. Eka Jaya).
  - "Schedule Data":   multi-stop circuits grouped by Route ID (e.g. Blue Water
                       Express, Gili Gili). Only OUTBOUND legs are turned into
                       sailings (boarding at a Bali port / Nusa Penida, getting
                       off at a later island or Lombok stop, up to the first
                       Bangsal stop). Return legs and mid-island boardings are
                       deliberately left out until an operator confirms them.
"""
import json, re, sys
from datetime import time, date, timedelta
import openpyxl

BALI = {"Sanur", "Padang Bai", "Serangan", "Kusamba", "Benoa / Nusa Dua"}
TURNAROUND = "Bangsal (Lombok)"
FLAG_TEXT = ("Published time differs from this operator's other departures on this route "
             "— confirm when booking.")
FLAG_CHECKIN = "Published check-in time looks unusual — confirm when booking."

# Season text in the sheet -> validity window (inclusive). Unknown text fails
# loudly so a new season label can't silently publish with no dates.
SEASONS = {
    "Nov '26 - March '27": ("2026-11-01", "2027-03-31"),
    "July - Sept '26": ("2026-07-01", "2026-09-30"),
    "Oct '26": ("2026-10-01", "2026-10-31"),
    "July '26 - Oct '26": ("2026-07-01", "2026-10-31"),
    "Confirmed on operator's website": (None, None),
    "Not stated on source": (None, None),
    "": (None, None),
}

def hhmm(v):
    if v in (None, ""):
        return None
    if isinstance(v, time):
        return f"{v.hour:02d}:{v.minute:02d}"
    s = str(v).strip()
    m = re.match(r"^(\d{1,2}):(\d{2})", s)
    if not m:
        raise ValueError(f"Unrecognised time value: {v!r}")
    return f"{int(m.group(1)):02d}:{m.group(2)}"

def iso(v):
    if isinstance(v, date):
        return v.isoformat()
    return str(v)[:10] if v else None

def season(text):
    text = (text or "").strip()
    if text not in SEASONS:
        raise SystemExit(f"Unknown season label {text!r} — add it to SEASONS in build_timetables.py")
    return text, *SEASONS[text]

def sid(*parts):
    return re.sub(r"[^a-z0-9]+", "-", "-".join(str(p) for p in parts).lower()).strip("-")


def merge_adjacent(sailings):
    """Identical sailings (same operator/ports/times/stops) in back-to-back seasons
    become ONE sailing with a combined window, so a service that simply carries on
    (e.g. Sept -> Oct) doesn't show as 'running now' AND 'starting soon'."""
    def key(x):
        return (x["operator"], x["from"], x["to"], x["departs"], x["arrives"], tuple(x["via"]), x["kind"], x["checkIn"])
    groups, order = {}, []
    for x in sailings:
        k = key(x)
        if k not in groups:
            groups[k] = []
            order.append(k)
        groups[k].append(x)
    out = []
    for k in order:
        items = groups[k]
        dated = sorted([i for i in items if i["validFrom"] and i["validTo"]], key=lambda i: i["validFrom"])
        undated = [i for i in items if not (i["validFrom"] and i["validTo"])]
        out.extend(undated)
        cur = None
        for i in dated:
            if cur and i["validFrom"] <= (date.fromisoformat(cur["validTo"]) + timedelta(days=1)).isoformat():
                cur["validTo"] = max(cur["validTo"], i["validTo"])
            else:
                if cur: out.append(cur)
                cur = dict(i)
        if cur: out.append(cur)
    return out

def main(path):
    wb = openpyxl.load_workbook(path)  # formulas ignored; we only need the typed-in cells
    sailings = []
    skipped = {}   # rows kept in the sheet but not published (no arrival time yet)
    held = {}      # rows whose arrival is an ESTIMATE: kept in the sheet, held back from the site
    parked = {}    # rows marked 'Not currently operating': kept in the sheet, ignored by the site

    # ---- Direct Sailings --------------------------------------------------
    ws = wb["Direct Sailings"]
    for r in range(2, ws.max_row + 1):
        op, frm, ci, dep, to, arr = (ws.cell(r, c).value for c in (1, 2, 3, 4, 5, 6))
        if not op or not frm or not to or str(op).lower().startswith("example"):
            continue
        if str(ws.cell(r, 23).value or "").startswith("Not currently operating"):
            parked[op] = parked.get(op, 0) + 1
            continue
        if not hhmm(dep) or not hhmm(arr):
            if hhmm(dep):
                skipped[op] = skipped.get(op, 0) + 1
            continue
        if str(ws.cell(r, 22).value or "").startswith("Estimated"):
            held[op] = held.get(op, 0) + 1
            continue
        s_text, v_from, v_to = season(ws.cell(r, 8).value)
        note = ws.cell(r, 14).value or ""
        flagged = bool(note) and not note.startswith(("Source page header", "INTERNAL:"))
        flag_text = FLAG_CHECKIN if note.startswith("Check-in") else FLAG_TEXT
        sailings.append({
            "id": sid(op, frm, to, hhmm(dep), v_from or "any"),
            "operator": op, "from": frm, "to": to,
            "checkIn": hhmm(ci), "departs": hhmm(dep), "arrives": hhmm(arr),
            "via": [], "kind": "direct",
            "season": None if s_text in ("Not stated on source", "") else s_text,
            "validFrom": v_from, "validTo": v_to,
            "source": ws.cell(r, 11).value, "verified": iso(ws.cell(r, 12).value),
            "flag": flag_text if flagged else None,
        })

    # ---- Schedule Data (multi-stop circuits) ------------------------------
    ws = wb["Schedule Data"]
    routes, order = {}, []
    pending = {}  # circuits marked PENDING CONFIRMATION: kept in the sheet, ignored by the site
    for r in range(2, ws.max_row + 1):
        rid, op, stop = (ws.cell(r, c).value for c in (1, 2, 3))
        if not rid or str(rid).upper().startswith("EXAMPLE") or not stop or not op:
            continue
        if rid not in routes:
            routes[rid] = {"operator": op, "stops": [], "season": ws.cell(r, 7).value,
                           "source": ws.cell(r, 10).value, "verified": iso(ws.cell(r, 11).value)}
            order.append(rid)
        routes[rid]["stops"].append((stop, hhmm(ws.cell(r, 4).value), hhmm(ws.cell(r, 5).value)))

    for rid in order:
        rt = routes[rid]
        if str(rt["season"] or "").startswith("PENDING CONFIRMATION"):
            pending[rt["operator"]] = pending.get(rt["operator"], 0) + 1
            continue
        s_text, v_from, v_to = season(rt["season"])
        stops = rt["stops"]
        # outbound = up to and including the first Bangsal stop
        end = next((i for i, s in enumerate(stops) if s[0] == TURNAROUND), len(stops) - 1)
        out = stops[: end + 1]
        for i, (p_i, _, dep_i) in enumerate(out):
            if dep_i is None or not (p_i in BALI or p_i == "Nusa Penida"):
                continue
            for j in range(i + 1, len(out)):
                p_j, arr_j, _ = out[j]
                if p_j in BALI or p_j == p_i or arr_j is None:
                    continue
                sailings.append({
                    "id": sid(rt["operator"], p_i, p_j, dep_i, rid),
                    "operator": rt["operator"], "from": p_i, "to": p_j,
                    "checkIn": None, "departs": dep_i, "arrives": arr_j,
                    "via": [s[0] for s in out[i + 1: j]], "kind": "circuit",
                    "season": None if s_text in ("Not stated on source", "Confirmed on operator's website", "") else s_text,
                    "validFrom": v_from, "validTo": v_to,
                    "source": rt["source"], "verified": rt["verified"], "flag": None,
                })

    sailings = merge_adjacent(sailings)

    ids = [s["id"] for s in sailings]
    dupes = {i for i in ids if ids.count(i) > 1}
    if dupes:
        raise SystemExit(f"Duplicate sailing ids: {sorted(dupes)[:5]}")

    updated = max(s["verified"] for s in sailings if s["verified"])
    lines = [
        "// AUTO-GENERATED by scripts/build_timetables.py from the operator research",
        "// spreadsheet. Do not edit by hand — change the spreadsheet and regenerate.",
        "// Port names match data/ports.js exactly. Times are 24h HH:MM, Bali time.",
        f'export const TIMETABLE_UPDATED = "{updated}";',
        "",
        "export const SAILINGS = [",
    ]
    for s in sailings:
        lines.append("  " + json.dumps(s, ensure_ascii=False) + ",")
    lines.append("];")
    out_path = "data/timetables.js"
    with open(out_path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines) + "\n")

    by_op = {}
    for s in sailings:
        by_op.setdefault((s["operator"], s["kind"]), 0)
        by_op[(s["operator"], s["kind"])] += 1
    print(f"Wrote {out_path}: {len(sailings)} sailings, updated {updated}")
    for (op, kind), n in by_op.items():
        print(f"  {op} ({kind}): {n}")
    for op, n in pending.items():
        print(f"  PENDING CONFIRMATION: {n} {op} circuit(s) held back — not published until confirmed.")
    for op, n in parked.items():
        print(f"  NOT OPERATING: {n} {op} rows are marked 'Not currently operating'; ignored by the site.")
    for op, n in held.items():
        print(f"  HELD BACK: {n} {op} rows have an ESTIMATED arrival time (see 'Arrival basis'); not published.")
    for op, n in skipped.items():
        print(f"  NOT PUBLISHED: {n} {op} rows have a departure but no arrival time (they stay in the spreadsheet).")

if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit(__doc__)
    main(sys.argv[1])
